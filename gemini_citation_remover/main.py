import re
from pathlib import Path
from docx import Document

input_path = Path("input")
output_path = Path("output")

# Create the output directory if it doesn't exist
output_path.mkdir(parents=True, exist_ok=True)

def clean_text(text: str) -> str:
  return re.sub(r"\[.*?\]", "", text)

# 1. Process .txt files
for file_path in input_path.glob("*.txt"):
  content = file_path.read_text(encoding="utf-8")
  cleaned = clean_text(content)

  destination = output_path / file_path.name
  destination.write_text(cleaned, encoding="utf-8")
  print(f"Saved cleaned file to: {destination}")

# 2. Process .docx files
for file_path in input_path.glob("*.docx"):
  if file_path.name.startswith("~$"):  # Skip temporary Word lock files
    continue

  doc = Document(file_path)

  # Clean body paragraphs
  for p in doc.paragraphs:
    if "[" in p.text:
      p.text = clean_text(p.text)

  # Clean text inside tables
  for table in doc.tables:
    for row in table.rows:
      for cell in row.cells:
        for p in cell.paragraphs:
          if "[" in p.text:
            p.text = clean_text(p.text)

  destination = output_path / file_path.name
  doc.save(destination)
  print(f"Saved cleaned file to: {destination}")